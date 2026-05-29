import os
import re
import uuid
from datetime import datetime
from io import BytesIO

import matplotlib.pyplot as plt
import pandas as pd
import plotly.express as px
import pymssql
import streamlit as st
from openpyxl.drawing.image import Image as OpenPyxlImage
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image as ReportLabImage,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


st.set_page_config(
    page_title="Well Depth Grapher",
    page_icon="📈",
    layout="wide"
)


WELL_OPTIONS = [
    "NONE",
    "Reeves",
    "Reeves A",
    "Park",
    "Park A",
    "Park B",
    "Woods",
    "Catawissa",
    "New",
    "Oakwood",
    "Ray",
    "Mt. Jefferson",
]


LINE_PATTERN = re.compile(
    r"^(?P<date>\d{4}/\d{2}/\d{2})\s+"
    r"(?P<time>\d{2}:\d{2}:\d{2}).*?"
    r"\bID\s+(?P<logger_id>\d+).*?"
    r"\bD\s+(?P<depth>-?\d+(?:\.\d+)?)\b"
)


def get_db_config():
    return {
        "server": os.getenv("DB_SERVER", "wwtp-sql"),
        "database": os.getenv("DB_NAME", "WWTP"),
        "user": os.getenv("DB_USER", "sa"),
        "password": os.getenv("DB_PASSWORD", ""),
        "port": int(os.getenv("DB_PORT", "1433")),
    }


def db_enabled():
    config = get_db_config()
    return bool(config["password"])


def get_connection():
    config = get_db_config()

    return pymssql.connect(
        server=config["server"],
        port=config["port"],
        user=config["user"],
        password=config["password"],
        database=config["database"],
        login_timeout=10,
        timeout=30,
    )


def ensure_tables_exist():
    sql = """
    IF NOT EXISTS (
        SELECT 1
        FROM sys.tables
        WHERE name = 'WellDepthChecks'
    )
    BEGIN
        CREATE TABLE dbo.WellDepthChecks (
            Id UNIQUEIDENTIFIER NOT NULL PRIMARY KEY,
            WellName NVARCHAR(100) NOT NULL,
            CheckDateTime DATETIME2 NOT NULL,
            SourceFiles NVARCHAR(MAX) NULL,
            Notes NVARCHAR(MAX) NULL,
            ReadingCount INT NOT NULL,
            MinDepth DECIMAL(10,2) NULL,
            MaxDepth DECIMAL(10,2) NULL,
            AvgDepth DECIMAL(10,2) NULL,
            CreatedAt DATETIME2 NOT NULL DEFAULT SYSUTCDATETIME()
        );
    END;

    IF NOT EXISTS (
        SELECT 1
        FROM sys.tables
        WHERE name = 'WellDepthReadings'
    )
    BEGIN
        CREATE TABLE dbo.WellDepthReadings (
            Id BIGINT IDENTITY(1,1) NOT NULL PRIMARY KEY,
            CheckId UNIQUEIDENTIFIER NOT NULL,
            ReadingDateTime DATETIME2 NOT NULL,
            LoggerId NVARCHAR(50) NULL,
            Depth DECIMAL(10,2) NOT NULL,
            SourceFile NVARCHAR(255) NULL,
            LineNumber INT NULL,
            CONSTRAINT FK_WellDepthReadings_WellDepthChecks
                FOREIGN KEY (CheckId)
                REFERENCES dbo.WellDepthChecks(Id)
                ON DELETE CASCADE
        );

        CREATE INDEX IX_WellDepthReadings_CheckId
            ON dbo.WellDepthReadings(CheckId);

        CREATE INDEX IX_WellDepthReadings_ReadingDateTime
            ON dbo.WellDepthReadings(ReadingDateTime);
    END;
    """

    with get_connection() as conn:
        with conn.cursor() as cursor:
            cursor.execute(sql)
        conn.commit()


def parse_txt_file(uploaded_file):
    rows = []

    text = uploaded_file.read().decode("utf-8", errors="replace")
    lines = text.splitlines()

    for line_number, line in enumerate(lines, start=1):
        match = LINE_PATTERN.search(line)

        if not match:
            continue

        date_text = match.group("date")
        time_text = match.group("time")
        logger_id = match.group("logger_id")
        depth = float(match.group("depth"))

        reading_datetime = pd.to_datetime(
            f"{date_text} {time_text}",
            format="%Y/%m/%d %H:%M:%S",
            errors="coerce",
        )

        if pd.isna(reading_datetime):
            continue

        rows.append(
            {
                "source_file": uploaded_file.name,
                "line_number": line_number,
                "logger_id": logger_id,
                "datetime": reading_datetime,
                "depth": depth,
            }
        )

    return pd.DataFrame(rows)


def save_depth_check(well_name, notes, df):
    ensure_tables_exist()

    check_id = str(uuid.uuid4())
    source_files = ", ".join(sorted(df["source_file"].dropna().unique()))

    check_datetime = df["datetime"].min().to_pydatetime()
    reading_count = int(len(df))
    min_depth = float(df["depth"].min())
    max_depth = float(df["depth"].max())
    avg_depth = float(df["depth"].mean())

    insert_check_sql = """
    INSERT INTO dbo.WellDepthChecks
    (
        Id,
        WellName,
        CheckDateTime,
        SourceFiles,
        Notes,
        ReadingCount,
        MinDepth,
        MaxDepth,
        AvgDepth
    )
    VALUES
    (
        %s, %s, %s, %s, %s, %s, %s, %s, %s
    );
    """

    insert_reading_sql = """
    INSERT INTO dbo.WellDepthReadings
    (
        CheckId,
        ReadingDateTime,
        LoggerId,
        Depth,
        SourceFile,
        LineNumber
    )
    VALUES
    (
        %s, %s, %s, %s, %s, %s
    );
    """

    reading_rows = []

    for _, row in df.iterrows():
        reading_rows.append(
            (
                check_id,
                row["datetime"].to_pydatetime(),
                str(row["logger_id"]),
                float(row["depth"]),
                str(row["source_file"]),
                int(row["line_number"]),
            )
        )

    with get_connection() as conn:
        with conn.cursor() as cursor:
            cursor.execute(
                insert_check_sql,
                (
                    check_id,
                    well_name,
                    check_datetime,
                    source_files,
                    notes,
                    reading_count,
                    min_depth,
                    max_depth,
                    avg_depth,
                ),
            )

            cursor.executemany(insert_reading_sql, reading_rows)

        conn.commit()

    return check_id


def load_saved_checks(well_name):
    ensure_tables_exist()

    sql = """
    SELECT
        Id,
        WellName,
        CheckDateTime,
        SourceFiles,
        Notes,
        ReadingCount,
        MinDepth,
        MaxDepth,
        AvgDepth,
        CreatedAt
    FROM dbo.WellDepthChecks
    WHERE WellName = %s
    ORDER BY CheckDateTime DESC;
    """

    with get_connection() as conn:
        df = pd.read_sql(sql, conn, params=[well_name])

    return df


def load_readings_for_checks(check_ids):
    if not check_ids:
        return pd.DataFrame()

    ensure_tables_exist()

    placeholders = ",".join(["%s"] * len(check_ids))

    sql = f"""
    SELECT
        c.Id AS CheckId,
        c.WellName,
        c.CheckDateTime,
        c.Notes,
        r.ReadingDateTime,
        r.LoggerId,
        r.Depth,
        r.SourceFile,
        r.LineNumber
    FROM dbo.WellDepthReadings r
    INNER JOIN dbo.WellDepthChecks c
        ON r.CheckId = c.Id
    WHERE c.Id IN ({placeholders})
    ORDER BY c.CheckDateTime, r.ReadingDateTime;
    """

    with get_connection() as conn:
        df = pd.read_sql(sql, conn, params=check_ids)

    return df


def build_check_label(row):
    check_date = pd.to_datetime(row["CheckDateTime"]).strftime("%Y-%m-%d %I:%M %p")
    return (
        f"{check_date} | "
        f"{int(row['ReadingCount'])} readings | "
        f"Min {row['MinDepth']:.2f} | "
        f"Max {row['MaxDepth']:.2f}"
    )


def safe_filename(value: str) -> str:
    value = value or "NONE"
    value = str(value).strip()
    value = re.sub(r"[^A-Za-z0-9_-]+", "_", value)
    return value.strip("_") or "NONE"


def build_export_dataframe(df: pd.DataFrame, well_name: str) -> pd.DataFrame:
    export_df = df.copy()

    export_df["datetime"] = pd.to_datetime(export_df["datetime"], errors="coerce")
    export_df["depth"] = pd.to_numeric(export_df["depth"], errors="coerce")

    export_df = export_df.dropna(subset=["datetime", "depth"])
    export_df = export_df.sort_values("datetime").reset_index(drop=True)

    export_df.insert(0, "well_name", well_name or "NONE")
    export_df["date"] = export_df["datetime"].dt.strftime("%m/%d/%Y")
    export_df["time"] = export_df["datetime"].dt.strftime("%H:%M:%S")

    return export_df[
        [
            "well_name",
            "date",
            "time",
            "datetime",
            "depth",
            "logger_id",
            "source_file",
            "line_number",
        ]
    ]


def create_depth_chart_image(export_df: pd.DataFrame, well_name: str) -> BytesIO:
    image_buffer = BytesIO()

    fig, ax = plt.subplots(figsize=(11, 5.5))

    ax.plot(
        export_df["datetime"],
        export_df["depth"],
        marker="o",
        linewidth=1.5,
        markersize=3,
    )

    ax.set_title(f"Well Depth Readings - {well_name or 'NONE'}")
    ax.set_xlabel("Date / Time")
    ax.set_ylabel("Well Depth")
    ax.grid(True, alpha=0.3)

    fig.autofmt_xdate()
    fig.tight_layout()

    fig.savefig(image_buffer, format="png", dpi=160)
    plt.close(fig)

    image_buffer.seek(0)
    return image_buffer


def create_excel_export(df: pd.DataFrame, well_name: str) -> BytesIO:
    export_df = build_export_dataframe(df, well_name)
    chart_image = create_depth_chart_image(export_df, well_name)

    output = BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        export_df.to_excel(writer, sheet_name="Depth Readings", index=False)

        workbook = writer.book
        worksheet = writer.sheets["Depth Readings"]

        worksheet.freeze_panes = "A2"

        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter

            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))

            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 32)

        chart_sheet = workbook.create_sheet("Chart")

        chart_sheet["A1"] = f"Well Depth Graph - {well_name or 'NONE'}"
        chart_sheet["A2"] = f"Exported: {datetime.now().strftime('%m/%d/%Y %H:%M:%S')}"
        chart_sheet["A3"] = f"Readings: {len(export_df):,}"

        if not export_df.empty:
            chart_sheet["A4"] = (
                f"Date Range: "
                f"{export_df['datetime'].min().strftime('%m/%d/%Y %H:%M:%S')} "
                f"through "
                f"{export_df['datetime'].max().strftime('%m/%d/%Y %H:%M:%S')}"
            )

        image_for_excel = OpenPyxlImage(chart_image)
        image_for_excel.anchor = "A6"
        chart_sheet.add_image(image_for_excel)

    output.seek(0)
    return output


def create_pdf_export(df: pd.DataFrame, well_name: str) -> BytesIO:
    export_df = build_export_dataframe(df, well_name)
    chart_image = create_depth_chart_image(export_df, well_name)

    output = BytesIO()

    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(letter),
        rightMargin=0.4 * inch,
        leftMargin=0.4 * inch,
        topMargin=0.4 * inch,
        bottomMargin=0.4 * inch,
    )

    styles = getSampleStyleSheet()
    story = []

    title = f"Well Depth Report - {well_name or 'NONE'}"
    exported_at = f"Exported: {datetime.now().strftime('%m/%d/%Y %H:%M:%S')}"

    if not export_df.empty:
        date_range = (
            f"Date Range: "
            f"{export_df['datetime'].min().strftime('%m/%d/%Y %H:%M:%S')} "
            f"through "
            f"{export_df['datetime'].max().strftime('%m/%d/%Y %H:%M:%S')}"
        )
    else:
        date_range = "Date Range: No readings"

    story.append(Paragraph(title, styles["Title"]))
    story.append(Paragraph(exported_at, styles["Normal"]))
    story.append(Paragraph(date_range, styles["Normal"]))
    story.append(Paragraph(f"Total Readings: {len(export_df):,}", styles["Normal"]))
    story.append(Spacer(1, 0.2 * inch))

    story.append(ReportLabImage(chart_image, width=9.8 * inch, height=4.7 * inch))
    story.append(Spacer(1, 0.25 * inch))

    table_data = [["Date", "Time", "Depth", "Logger ID", "Source File"]]

    for _, row in export_df.iterrows():
        table_data.append(
            [
                row["date"],
                row["time"],
                f"{row['depth']:.2f}",
                str(row["logger_id"]),
                str(row["source_file"]),
            ]
        )

    table = Table(table_data, repeatRows=1)

    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#eff6ff")],
                ),
                ("ALIGN", (2, 1), (2, -1), "RIGHT"),
            ]
        )
    )

    story.append(table)
    doc.build(story)

    output.seek(0)
    return output


st.title("Well Depth Grapher")

st.write(
    "Upload TXT logger files, graph well depth over time, and optionally save named well checks to the database."
)

tab_upload, tab_history = st.tabs(["Upload / Graph", "History / Compare"])


with tab_upload:
    st.subheader("Upload New TXT File")

    selected_well = st.selectbox(
        "Well Name",
        WELL_OPTIONS,
        index=0,
        help="Choose NONE if this is not one of the Town wells or if you do not want to save the results.",
    )

    notes = st.text_area(
        "Notes",
        placeholder="Optional notes, such as drought conditions, rainfall, pump test notes, or anything unusual.",
    )

    uploaded_files = st.file_uploader(
        "Upload TXT file(s)",
        type=["txt"],
        accept_multiple_files=True,
    )

    if uploaded_files:
        all_data = []

        for uploaded_file in uploaded_files:
            df_file = parse_txt_file(uploaded_file)
            all_data.append(df_file)

        df = pd.concat(all_data, ignore_index=True)

        if df.empty:
            st.error("No depth readings were found. Make sure the file has lines with a D value.")
        else:
            df = df.dropna(subset=["datetime", "depth"])
            df = df.sort_values("datetime")

            st.success(f"Parsed {len(df):,} depth readings.")

            col1, col2, col3, col4, col5 = st.columns(5)

            col1.metric("Readings", f"{len(df):,}")
            col2.metric("Minimum Depth", f"{df['depth'].min():.2f}")
            col3.metric("Maximum Depth", f"{df['depth'].max():.2f}")
            col4.metric("Average Depth", f"{df['depth'].mean():.2f}")
            col5.metric("Files", f"{len(uploaded_files):,}")

            st.subheader("Depth Over Time")

            chart_title = (
                "Uploaded Well Depth Readings"
                if selected_well == "NONE"
                else f"{selected_well} Well Depth Readings"
            )

            fig = px.line(
                df,
                x="datetime",
                y="depth",
                color="logger_id",
                title=chart_title,
                labels={
                    "datetime": "Date / Time",
                    "depth": "Well Depth",
                    "logger_id": "Logger ID",
                },
            )

            fig.update_layout(
                hovermode="x unified",
                yaxis_title="Well Depth",
                xaxis_title="Date / Time",
            )

            st.plotly_chart(fig, use_container_width=True)

            st.subheader("Data Table")
            st.dataframe(df, use_container_width=True)

            st.subheader("Download Report")

            export_well_name = selected_well if selected_well else "NONE"
            safe_well = safe_filename(export_well_name)
            file_stamp = datetime.now().strftime("%Y%m%d_%H%M%S")

            csv = df.to_csv(index=False).encode("utf-8")

            try:
                excel_file = create_excel_export(df, export_well_name)
                pdf_file = create_pdf_export(df, export_well_name)

                col_download_1, col_download_2, col_download_3 = st.columns(3)

                with col_download_1:
                    st.download_button(
                        "Download Parsed CSV",
                        data=csv,
                        file_name=f"well_depth_{safe_well}_{file_stamp}.csv",
                        mime="text/csv",
                    )

                with col_download_2:
                    st.download_button(
                        "Download Excel Report",
                        data=excel_file,
                        file_name=f"well_depth_{safe_well}_{file_stamp}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )

                with col_download_3:
                    st.download_button(
                        "Download PDF Report",
                        data=pdf_file,
                        file_name=f"well_depth_{safe_well}_{file_stamp}.pdf",
                        mime="application/pdf",
                    )

            except Exception as export_error:
                st.error(f"Could not create export files: {export_error}")

            if selected_well == "NONE":
                st.info("Well Name is set to NONE, so this upload will not be saved to the database.")
            else:
                if not db_enabled():
                    st.error(
                        "Database password is not configured. Add DB_PASSWORD in Portainer environment variables."
                    )
                else:
                    if st.button("Save This Depth Check to Database"):
                        try:
                            check_id = save_depth_check(selected_well, notes, df)
                            st.success(
                                f"Saved {len(df):,} readings for {selected_well}."
                            )
                        except Exception as ex:
                            st.error(f"Save failed: {ex}")


with tab_history:
    st.subheader("History / Compare Previous Checks")

    history_well = st.selectbox(
        "Choose Well for History",
        [well for well in WELL_OPTIONS if well != "NONE"],
    )

    if not db_enabled():
        st.warning("Database is not configured yet. Add DB_PASSWORD in Portainer.")
    else:
        if st.button("Load History"):
            st.session_state["history_loaded"] = True

        if st.session_state.get("history_loaded"):
            try:
                checks_df = load_saved_checks(history_well)

                if checks_df.empty:
                    st.info(f"No saved depth checks found for {history_well}.")
                else:
                    checks_df["Label"] = checks_df.apply(build_check_label, axis=1)

                    st.write(f"Saved checks for **{history_well}**:")

                    display_df = checks_df[
                        [
                            "CheckDateTime",
                            "ReadingCount",
                            "MinDepth",
                            "MaxDepth",
                            "AvgDepth",
                            "SourceFiles",
                            "Notes",
                        ]
                    ].copy()

                    st.dataframe(display_df, use_container_width=True)

                    selected_labels = st.multiselect(
                        "Select checks to compare",
                        checks_df["Label"].tolist(),
                        default=checks_df["Label"].head(2).tolist(),
                    )

                    selected_check_ids = checks_df[
                        checks_df["Label"].isin(selected_labels)
                    ]["Id"].astype(str).tolist()

                    if selected_check_ids:
                        readings_df = load_readings_for_checks(selected_check_ids)

                        if readings_df.empty:
                            st.warning("No readings found for the selected checks.")
                        else:
                            readings_df["CheckLabel"] = readings_df.apply(
                                lambda row: pd.to_datetime(row["CheckDateTime"]).strftime(
                                    "%Y-%m-%d %I:%M %p"
                                ),
                                axis=1,
                            )

                            st.subheader("Comparison Chart")

                            fig_history = px.line(
                                readings_df,
                                x="ReadingDateTime",
                                y="Depth",
                                color="CheckLabel",
                                title=f"{history_well} Depth Check Comparison",
                                labels={
                                    "ReadingDateTime": "Date / Time",
                                    "Depth": "Well Depth",
                                    "CheckLabel": "Depth Check",
                                },
                            )

                            fig_history.update_layout(
                                hovermode="x unified",
                                yaxis_title="Well Depth",
                                xaxis_title="Date / Time",
                            )

                            st.plotly_chart(fig_history, use_container_width=True)

                            st.subheader("Comparison Summary")

                            summary_df = (
                                readings_df.groupby("CheckLabel")
                                .agg(
                                    Readings=("Depth", "count"),
                                    MinDepth=("Depth", "min"),
                                    MaxDepth=("Depth", "max"),
                                    AvgDepth=("Depth", "mean"),
                                    StartTime=("ReadingDateTime", "min"),
                                    EndTime=("ReadingDateTime", "max"),
                                )
                                .reset_index()
                            )

                            st.dataframe(summary_df, use_container_width=True)

                            st.subheader("Raw Historical Readings")
                            st.dataframe(readings_df, use_container_width=True)

            except Exception as ex:
                st.error(f"Could not load history: {ex}")
